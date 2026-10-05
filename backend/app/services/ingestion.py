"""Bounded CSV/XLSX readers. No macros, formula evaluation or executable deserialization."""
from io import BytesIO
from pathlib import Path
from zipfile import BadZipFile, ZipFile
import csv

import pandas as pd
from openpyxl import load_workbook

MAX_ROWS = 1_000_000
MAX_COLUMNS = 200


def workbook_sheets(content: bytes) -> list[str]:
    try:
        with ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > 10000 or sum(e.file_size for e in entries) > 300 * 1024 * 1024:
                raise ValueError("Workbook expands beyond the 300 MB safety limit.")
        book = load_workbook(BytesIO(content), read_only=True, data_only=True, keep_links=False)
        try:
            return book.sheetnames
        finally:
            book.close()
    except (BadZipFile, OSError, KeyError) as exc:
        raise ValueError("Invalid XLSX workbook.") from exc


def read_table(content: bytes, filename: str, sheet: str | None = None) -> pd.DataFrame:
    suffix = Path(filename).suffix.lower()
    if suffix == ".csv":
        last_error = None
        for encoding in ["utf-8-sig", "cp1251"]:
            try:
                prefix = content[:65536].decode(encoding, errors="ignore")
                # Avoid a cut-off final row misleading the delimiter sniffer.
                prefix = prefix.rsplit("\n", 1)[0] or prefix
                try:
                    separator = csv.Sniffer().sniff(prefix, delimiters=",;\t|").delimiter
                except csv.Error:
                    separator = ","
                frame = pd.read_csv(BytesIO(content), encoding=encoding, sep=separator, dtype=object, nrows=MAX_ROWS + 1)
                break
            except UnicodeDecodeError as exc:
                last_error = exc
        else:
            raise ValueError("CSV encoding must be UTF-8 or Windows-1251.") from last_error
    elif suffix == ".xlsx":
        names = workbook_sheets(content)
        if sheet not in names:
            raise ValueError("Choose a valid workbook sheet.")
        frame = pd.read_excel(BytesIO(content), sheet_name=sheet, engine="openpyxl", nrows=MAX_ROWS + 1)
    else:
        raise ValueError("Only .csv and .xlsx files are supported.")
    if len(frame) > MAX_ROWS or len(frame.columns) > MAX_COLUMNS:
        raise ValueError("Dataset exceeds 1,000,000 rows or 200 columns. Split the file before uploading.")
    if not len(frame.columns):
        raise ValueError("No header columns found.")
    names, seen = [], set()
    for column in frame.columns:
        base = str(column).strip() or "Unnamed"
        name, number = base, 2
        while name in seen:
            name, number = f"{base}_{number}", number + 1
        names.append(name)
        seen.add(name)
    frame.columns = names
    for name in frame.select_dtypes(include=["object", "string"]):
        frame[name] = frame[name].map(lambda v: v.strip() if isinstance(v, str) else v)
        frame[name] = frame[name].replace("", None)
    return frame
